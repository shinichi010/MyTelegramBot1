import io
import os
import html
import asyncio
import logging
import threading
import time
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import httpx
from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputMediaPhoto,
    InputMediaDocument,
)
from telegram.error import BadRequest
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)
from motor.motor_asyncio import AsyncIOMotorClient

try:
    from openpyxl import Workbook
except ImportError:  # التصدير يتعطل فقط إذا المكتبة غير مثبتة
    Workbook = None

# ============================================================
# الإعدادات (المتغيرات البيئية)
# ============================================================
BOT_TOKEN = os.getenv("BOT_TOKEN")
MONGO_URI = os.getenv("MONGO_URI")
CHANNEL_ID_RAW = os.getenv("CHANNEL_ID")  # آيدي القناة، مثال: -1001234567890

try:
    CHANNEL_ID = int(CHANNEL_ID_RAW) if CHANNEL_ID_RAW else None
except ValueError:
    CHANNEL_ID = None

OWNER_USERNAME = "snh_1"   # المالك الأساسي (بدون @)
MIN_NAME_WORDS = 2         # أقل عدد كلمات للاسم

PORT = int(os.getenv("PORT", "10000"))  # Render يعطي المنفذ تلقائياً
PING_URL = os.getenv("PING_URL") or os.getenv("RENDER_EXTERNAL_URL")

BAGHDAD = timezone(timedelta(hours=3))  # توقيت العراق

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)
logging.getLogger("httpx").setLevel(logging.WARNING)  # حتى ما يطبع التوكن بالـ Logs

# ============================================================
# قاعدة البيانات
# ============================================================
client = AsyncIOMotorClient(MONGO_URI)
db = client.graduation_bot
# الكولكشنات:
#   config         -> الإعدادات ورسائل البوت
#   users          -> المستخدمين (بياناتهم + حالة الخطوة + المسودة)
#   students       -> منشورات الطلاب (الاسم + الصورة + رقم المنشور بالقناة)
#   pending_edits  -> طلبات التعديل المعلقة
#   admins         -> المشرفين وصلاحياتهم
#   banned         -> المحظورين

# ============================================================
# الصلاحيات
# ============================================================
PERMS = {
    "stats": "📊 الإحصائيات",
    "broadcast": "📢 الإذاعة والتذكير",
    "messages": "📝 تعديل رسائل البوت",
    "toggle": "⏯ إيقاف/تشغيل الاستقبال والتعديل والإغلاق التلقائي",
    "ban": "🚫 الحظر وفك الحظر",
    "edits": "✅ الموافقة على طلبات التعديل",
    "ping": "🏓 التحكم بالبينك",
    "export": "📤 تصدير القائمة",
    "delete": "🗑 حذف الطلبات",
    "report": "📋 التقرير اليومي",
}

# ============================================================
# رسائل البوت (قابلة للتعديل من لوحة الإدارة)
# ============================================================
DEFAULT_MESSAGES = {
    "start_msg": "هلا بيك 👋\nهذا بوت استلام صور التخرج.\n(للمساعدة ارسل /help)",
    "ask_name_msg": "اكتب اسمك الثلاثي من فضلك ✍️",
    "bad_name_msg": "اكتب اسمك كامل (كلمتين على الأقل) وبنص فقط.",
    "ask_photo_msg": "تمام ✅\nهسه ارسل صورتك.",
    "wrong_type_msg": "ارسل صورة فقط 🖼 (صورة عادية أو ملف بصيغة صورة).",
    "preview_msg": "هذا اسمك وصورتك، تأكد منهم قبل الإرسال 👇",
    "confirm_hint_msg": "اضغط ✅ تأكيد أو 🔄 إعادة من الرسالة اللي فوق.",
    "success_msg": "تم استلام اسمك وصورتك بنجاح ✅",
    "edit_note_msg": "إذا تريد تعدل اسمك أو صورتك اكتب: تعديل",
    "already_msg": "تم استلام طلبك مسبقاً ✅",
    "edit_ask_name_msg": "اكتب الاسم الثلاثي الجديد ✍️\n(للإلغاء اكتب: الغاء)",
    "edit_ask_photo_msg": "هسه ارسل الصورة الجديدة 🖼\n(للإلغاء اكتب: الغاء)",
    "edit_sent_msg": "تم إرسال طلب التعديل، راح توصلك رسالة بالنتيجة ⏳",
    "edit_pending_msg": "عندك طلب تعديل قيد المراجعة، انتظر النتيجة ⏳",
    "edit_approved_msg": "تمت الموافقة على طلب التعديل ✅",
    "edit_rejected_msg": "تم رفض طلب التعديل ❌",
    "edit_disabled_msg": "التعديل متوقف حالياً.",
    "deleted_msg": "تم حذف طلبك السابق، تكدر ترسل اسمك وصورتك من جديد.",
    "maintenance_msg": "نعتذر، استلام الطلبات متوقف حالياً.",
    "help_msg": f"إذا واجهتك مشكلة تواصل مع المطور: @{OWNER_USERNAME}",
    "not_admin_msg": "عذراً، هذا الأمر للمشرفين فقط.",
    "error_msg": "صار خطأ، حاول مرة ثانية بعد شوية.",
}

MSG_LABELS = {
    "start_msg": "👋 الترحيب",
    "ask_name_msg": "✍️ طلب الاسم",
    "bad_name_msg": "⚠️ اسم غير صحيح",
    "ask_photo_msg": "🖼 طلب الصورة",
    "wrong_type_msg": "⚠️ مو صورة",
    "preview_msg": "👁 المعاينة",
    "confirm_hint_msg": "⏳ تنبيه التأكيد",
    "success_msg": "✅ الاستلام",
    "edit_note_msg": "ℹ️ ملاحظة التعديل",
    "already_msg": "🔁 مسجل مسبقاً",
    "edit_ask_name_msg": "✏️ تعديل: الاسم",
    "edit_ask_photo_msg": "✏️ تعديل: الصورة",
    "edit_sent_msg": "📤 تعديل: أُرسل",
    "edit_pending_msg": "⏳ تعديل: قيد المراجعة",
    "edit_approved_msg": "✅ تعديل: القبول",
    "edit_rejected_msg": "❌ تعديل: الرفض",
    "edit_disabled_msg": "🚫 التعديل متوقف",
    "deleted_msg": "🗑 حذف الطلب",
    "maintenance_msg": "⏸ الاستقبال متوقف",
    "help_msg": "❓ المساعدة",
    "not_admin_msg": "⛔ /admin لغير المشرف",
    "error_msg": "⚠️ رسالة الخطأ",
}

DEFAULT_CONFIG = {
    "maintenance_mode": False,
    "edit_enabled": True,
    "total_received": 0,
    "ping_enabled": True,
    "ping_interval": 10,  # بالدقائق
    "close_at": None,     # موعد الإغلاق التلقائي (UTC)
    "report_enabled": False,   # التقرير اليومي للقناة
    "report_time": "22:00",    # وقت التقرير (توقيت العراق)
    "report_last_date": None,
    **DEFAULT_MESSAGES,
}


# ============================================================
# دوال مساعدة عامة
# ============================================================
def btn(text, data):
    return InlineKeyboardButton(text, callback_data=data)


async def get_config():
    stored = await db.config.find_one({"_id": "main_config"}) or {}
    return {**DEFAULT_CONFIG, **stored}


async def set_config(key, value):
    await db.config.update_one({"_id": "main_config"}, {"$set": {key: value}}, upsert=True)


async def inc_received():
    await db.config.update_one({"_id": "main_config"}, {"$inc": {"total_received": 1}}, upsert=True)


def close_time(conf):
    ca = conf.get("close_at")
    if ca and ca.tzinfo is None:
        ca = ca.replace(tzinfo=timezone.utc)
    return ca


def is_closed(conf):
    """الاستقبال مغلق: يدوياً (صيانة) أو بسبب موعد الإغلاق التلقائي."""
    if conf.get("maintenance_mode"):
        return True
    ca = close_time(conf)
    return bool(ca and datetime.now(timezone.utc) >= ca)


def is_owner(user):
    return bool(user and user.username and user.username.lower() == OWNER_USERNAME.lower())


async def is_admin(user):
    if is_owner(user):
        return True
    return bool(await db.admins.find_one({"user_id": user.id}))


async def perms_of(user):
    """ترجع مجموعة الصلاحيات الفعلية للمستخدم."""
    if is_owner(user):
        return set(PERMS.keys())
    doc = await db.admins.find_one({"user_id": user.id})
    if not doc:
        return set()
    return {k for k, v in doc.get("perms", {}).items() if v}


async def has_perm(user, perm):
    return perm in await perms_of(user)


async def is_banned(user_id):
    return bool(await db.banned.find_one({"user_id": user_id}))


async def set_state(user_id, state, mode=None, tmp_name=None, draft=None):
    await db.users.update_one(
        {"user_id": user_id},
        {"$set": {"state": state, "mode": mode, "tmp_name": tmp_name, "draft": draft}},
        upsert=True,
    )


def with_edit_note(conf, key):
    text = conf[key]
    if conf.get("edit_enabled"):
        text += "\n\n" + conf["edit_note_msg"]
    return text


CANCEL_WORDS = {"الغاء", "إلغاء"}  # كلمة الإلغاء (تنقبل بالهمزة وبدونها)


def is_cancel(text):
    return (text or "").strip() in CANCEL_WORDS


def valid_name(text):
    return len(text.split()) >= MIN_NAME_WORDS and len(text) <= 100


def extract_image(msg):
    if msg.photo:
        return msg.photo[-1].file_id, "photo"
    d = msg.document
    if d and d.mime_type and d.mime_type.startswith("image/"):
        return d.file_id, "document"
    return None, None


async def say(context, user_id, text):
    await context.bot.send_message(user_id, text)


async def safe_edit(query, text, markup=None):
    try:
        await query.edit_message_text(text, reply_markup=markup, parse_mode="HTML")
    except BadRequest as e:
        if "not modified" not in str(e).lower():
            logger.warning(f"safe_edit: {e}")


def who_is(user):
    s = html.escape(user.full_name)
    if user.username:
        s += f" (@{user.username})"
    return s


# ============================================================
# المستخدمين + إشعار المالك عند أول ستارت
# ============================================================
async def notify_owner_new_user(user, context):
    photo_id = None
    try:
        photos = await context.bot.get_user_profile_photos(user.id, limit=1)
        if photos.total_count > 0:
            photo_id = photos.photos[0][-1].file_id
    except Exception as e:
        logger.warning(f"profile photo: {e}")

    if photo_id:
        await db.users.update_one({"user_id": user.id}, {"$set": {"profile_photo": photo_id}})

    conf = await get_config()
    owner_id = conf.get("owner_id")
    if not owner_id:
        return

    username = f"@{user.username}" if user.username else "لا يوجد"
    text = (
        "🆕 <b>مستخدم جديد ضغط ستارت</b>\n"
        f"👤 الاسم: {html.escape(user.full_name)}\n"
        f"🔗 اليوزر: {username}\n"
        f"🆔 الآيدي: <code>{user.id}</code>"
    )
    try:
        if photo_id:
            await context.bot.send_photo(owner_id, photo_id, caption=text, parse_mode="HTML")
        else:
            await context.bot.send_message(owner_id, text + "\n🖼 الصورة: لا توجد/مخفية", parse_mode="HTML")
    except Exception as e:
        logger.error(f"owner notify failed: {e}")


async def touch_user(user, context):
    """يسجل/يحدث المستخدم. إذا أول مرة يدخل يبلغ المالك. ترجع True إذا جديد."""
    now = datetime.now(timezone.utc)
    res = await db.users.update_one(
        {"user_id": user.id},
        {
            "$set": {"full_name": user.full_name, "username": user.username},
            "$setOnInsert": {"joined_at": now},
        },
        upsert=True,
    )
    # تحديث اسم المشرف (إذا كان مشرف) حتى يظهر بالقائمة
    await db.admins.update_one(
        {"user_id": user.id},
        {"$set": {"name": user.full_name, "username": user.username}},
    )
    if is_owner(user):
        await set_config("owner_id", user.id)

    is_new = res.upserted_id is not None
    if is_new and not is_owner(user):
        await notify_owner_new_user(user, context)
    return is_new


async def notify_owner_error(context, text):
    conf = await get_config()
    owner_id = conf.get("owner_id")
    if owner_id:
        try:
            await context.bot.send_message(owner_id, text)
        except Exception:
            pass


# ============================================================
# النشر بالقناة
# ============================================================
def request_keyboard(uid):
    return InlineKeyboardMarkup([
        [btn("✅ قبول", f"ok_{uid}"), btn("❌ رفض", f"no_{uid}")],
    ])


async def send_file(bot, chat_id, file_id, ftype, caption, markup):
    if ftype == "photo":
        return await bot.send_photo(chat_id, file_id, caption=caption, parse_mode="HTML", reply_markup=markup)
    return await bot.send_document(chat_id, file_id, caption=caption, parse_mode="HTML", reply_markup=markup)


async def send_post(bot, uid, name, file_id, ftype):
    return await send_file(bot, CHANNEL_ID, file_id, ftype, f"👤 {html.escape(name)}", None)


def info_text(full_name, username, uid):
    uname = f"@{html.escape(username)}" if username else "لا يوجد"
    return (
        f"👤 اسم الحساب: {html.escape(full_name or '-')}\n"
        f"🔗 اليوزر: {uname}\n"
        f"🆔 الآيدي: <code>{uid}</code>"
    )


async def send_info(bot, uid, full_name, username, reply_to):
    """يرسل رسالة معلومات الحساب كرد على منشور الطالب بالقناة. يرجع رقمها أو None."""
    try:
        m = await bot.send_message(
            CHANNEL_ID,
            info_text(full_name, username, uid),
            parse_mode="HTML",
            reply_to_message_id=reply_to,
            allow_sending_without_reply=True,
        )
        return m.message_id
    except Exception as e:
        logger.warning(f"send_info failed: {e}")
        return None


async def delete_post_msgs(bot, d):
    """يحذف منشور الطالب ورسالة معلوماته من القناة. يرجع True إذا انحذف المنشور الأساسي."""
    ok = True
    try:
        await bot.delete_message(CHANNEL_ID, d["channel_msg_id"])
    except Exception as e:
        ok = False
        logger.warning(f"delete channel post failed: {e}")
    if d.get("info_msg_id"):
        try:
            await bot.delete_message(CHANNEL_ID, d["info_msg_id"])
        except Exception as e:
            logger.warning(f"delete info msg failed: {e}")
    return ok


async def send_request(bot, uid, old_name, new_name, file_id, ftype):
    caption = (
        "🔄 <b>طلب تعديل</b>\n"
        f"👤 الاسم القديم: {html.escape(old_name)}\n"
        f"✏️ الاسم الجديد: {html.escape(new_name)}"
    )
    return await send_file(bot, CHANNEL_ID, file_id, ftype, caption, request_keyboard(uid))


async def replace_post(bot, uid, student, new_name, file_id, ftype):
    """يعدل منشور الطالب القديم بالقناة. إذا فشل، ينشر جديد ويحذف القديم.
    يرجع (رقم المنشور، رقم رسالة المعلومات)."""
    caption = f"👤 {html.escape(new_name)}"
    old_id = (student or {}).get("channel_msg_id")
    old_info = (student or {}).get("info_msg_id")
    u = await db.users.find_one({"user_id": uid}) or {}
    full_name, username = u.get("full_name"), u.get("username")
    if old_id:
        if ftype == "photo":
            media = InputMediaPhoto(file_id, caption=caption, parse_mode="HTML")
        else:
            media = InputMediaDocument(file_id, caption=caption, parse_mode="HTML")
        try:
            await bot.edit_message_media(chat_id=CHANNEL_ID, message_id=old_id, media=media)
            return old_id, await refresh_info(bot, uid, old_id, old_info, full_name, username)
        except Exception as e:
            if "not modified" in str(e).lower():
                return old_id, await refresh_info(bot, uid, old_id, old_info, full_name, username)
            logger.warning(f"edit_message_media failed, reposting: {e}")

    sent = await send_post(bot, uid, new_name, file_id, ftype)
    if old_id:
        try:
            await bot.delete_message(CHANNEL_ID, old_id)
        except Exception:
            pass
    if old_info:
        try:
            await bot.delete_message(CHANNEL_ID, old_info)
        except Exception:
            pass
    info_id = await send_info(bot, uid, full_name, username, sent.message_id)
    return sent.message_id, info_id


async def refresh_info(bot, uid, post_id, info_id, full_name, username):
    """يحدّث رسالة المعلومات (أو ينشئها إذا ما موجودة). يرجع رقمها."""
    if info_id:
        try:
            await bot.edit_message_text(
                info_text(full_name, username, uid),
                chat_id=CHANNEL_ID, message_id=info_id, parse_mode="HTML",
            )
            return info_id
        except Exception as e:
            if "not modified" in str(e).lower():
                return info_id
            logger.warning(f"edit info failed, resending: {e}")
            try:
                await bot.delete_message(CHANNEL_ID, info_id)
            except Exception:
                pass
    return await send_info(bot, uid, full_name, username, post_id)


# ============================================================
# الأوامر الأساسية
# ============================================================
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    msg = update.message
    if await is_banned(user.id):
        return

    await touch_user(user, context)
    conf = await get_config()
    admin = await is_admin(user)

    if is_closed(conf) and not admin:
        await msg.reply_text(conf["maintenance_msg"])
        return

    if admin:
        await msg.reply_text("أهلاً بيك 👋\nلوحة التحكم: /admin")

    u = await db.users.find_one({"user_id": user.id}) or {}
    # المشرفين يكدرون يرسلون أكثر من مرة، فما نعتبرهم "مسجلين"
    student = None if admin else await db.students.find_one({"user_id": user.id})
    state = u.get("state")
    mode = u.get("mode") or "new"

    if state == "confirm":
        await msg.reply_text(conf["confirm_hint_msg"])
    elif state == "name":
        await msg.reply_text(conf["edit_ask_name_msg"] if mode == "edit" else conf["ask_name_msg"])
    elif state == "photo":
        await msg.reply_text(conf["edit_ask_photo_msg"] if mode == "edit" else conf["ask_photo_msg"])
    elif student:
        await msg.reply_text(with_edit_note(conf, "already_msg"))
    else:
        await set_state(user.id, "name", "new")
        await msg.reply_text(conf["start_msg"])
        await msg.reply_text(conf["ask_name_msg"])


async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if await is_banned(user.id):
        return
    await touch_user(user, context)
    conf = await get_config()
    await update.message.reply_text(conf["help_msg"])


# ============================================================
# الحظر / المشرفين / الحذف / التصدير (أوامر)
# ============================================================
async def ban_user(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await has_perm(update.effective_user, "ban"):
        return
    if not context.args:
        await update.message.reply_text("مثال: /ban 123456789")
        return
    try:
        target = int(context.args[0])
    except ValueError:
        await update.message.reply_text("❌ الآيدي غير صحيح.")
        return
    if await db.admins.find_one({"user_id": target}):
        await update.message.reply_text("❌ ما يصير تحظر مشرف.")
        return
    await db.banned.update_one({"user_id": target}, {"$set": {"user_id": target}}, upsert=True)
    await update.message.reply_text(f"🚫 تم حظر {target}.")


async def unban_user(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await has_perm(update.effective_user, "ban"):
        return
    if not context.args:
        await update.message.reply_text("مثال: /unban 123456789")
        return
    try:
        target = int(context.args[0])
    except ValueError:
        await update.message.reply_text("❌ الآيدي غير صحيح.")
        return
    await db.banned.delete_one({"user_id": target})
    await update.message.reply_text(f"✅ تم فك الحظر عن {target}.")


async def add_admin(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update.effective_user):
        return
    if not context.args:
        await update.message.reply_text("مثال: /add_admin 123456789")
        return
    try:
        admin_id = int(context.args[0])
    except ValueError:
        await update.message.reply_text("❌ الآيدي غير صحيح.")
        return
    await db.admins.update_one(
        {"user_id": admin_id},
        {"$setOnInsert": {"perms": {k: False for k in PERMS}}},
        upsert=True,
    )
    await update.message.reply_text(
        f"✅ تمت إضافة {admin_id} كمشرف (بدون صلاحيات حالياً).",
        reply_markup=InlineKeyboardMarkup([[btn("⚙️ ضبط الصلاحيات", f"adm_{admin_id}")]]),
    )


async def rem_admin(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_owner(update.effective_user):
        return
    if not context.args:
        await update.message.reply_text("مثال: /rem_admin 123456789")
        return
    try:
        admin_id = int(context.args[0])
    except ValueError:
        await update.message.reply_text("❌ الآيدي غير صحيح.")
        return
    await db.admins.delete_one({"user_id": admin_id})
    await update.message.reply_text(f"✅ تمت إزالة {admin_id} من المشرفين.")


async def delete_request(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/delete ID — يحذف منشور الطالب من القناة ومن القاعدة ويخليه يرسل من جديد."""
    if not await has_perm(update.effective_user, "delete"):
        return
    if not context.args:
        await update.message.reply_text("مثال: /delete 123456789")
        return
    try:
        target = int(context.args[0])
    except ValueError:
        await update.message.reply_text("❌ الآيدي غير صحيح.")
        return

    docs = [d async for d in db.students.find({"user_id": target})]
    pend = await db.pending_edits.find_one_and_delete({"user_id": target})
    if not docs and not pend:
        await update.message.reply_text("ما لكيت طلب لهذا الآيدي.")
        return

    failed = 0
    for d in docs:
        if not await delete_post_msgs(context.bot, d):
            failed += 1
    if pend and pend.get("request_msg_id"):
        try:
            await context.bot.delete_message(CHANNEL_ID, pend["request_msg_id"])
        except Exception:
            pass

    await db.students.delete_many({"user_id": target})
    # نرجع حالته للصفر (بدون ما ننشئ سجل لمستخدم ما دخل البوت)
    await db.users.update_one(
        {"user_id": target},
        {"$set": {"state": None, "mode": None, "tmp_name": None, "draft": None}},
    )

    conf = await get_config()
    try:
        await context.bot.send_message(target, conf["deleted_msg"])
    except Exception:
        pass

    text = f"✅ تم حذف طلب {target} من القاعدة."
    if failed:
        text += f"\n⚠️ فشل حذف {failed} منشور من القناة، احذفه يدوياً."
    await update.message.reply_text(text)


def parse_dt(text):
    """يقرأ تاريخ (مع وقت اختياري) بتوقيت العراق."""
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=BAGHDAD)
        except ValueError:
            pass
    return None


async def delete_before_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/delete_before التاريخ — يحذف كل المنشورات اللي انرسلت قبل هذا التاريخ (بعد تأكيد)."""
    if not await has_perm(update.effective_user, "delete"):
        return
    text = " ".join(context.args).strip()
    dt = parse_dt(text) if text else None
    if not dt:
        await update.message.reply_text(
            "مثال: /delete_before 2026-10-20\nأو مع الوقت: /delete_before 2026-10-20 14:30\n(توقيت العراق)"
        )
        return
    if dt > datetime.now(timezone.utc):
        await update.message.reply_text("❌ التاريخ لازم يكون بالماضي.")
        return

    cutoff = dt.astimezone(timezone.utc)
    count = await db.students.count_documents({"submitted_at": {"$lt": cutoff}})
    if not count:
        await update.message.reply_text("ما كو منشورات قبل هذا التاريخ.")
        return

    shown = dt.strftime("%Y-%m-%d %H:%M")
    await update.message.reply_text(
        f"⚠️ راح يتم حذف <b>{count}</b> منشور (من القناة ومن القاعدة) انرسلوا قبل <b>{shown}</b> (توقيت العراق).\n"
        "هذا الإجراء ما يرجع.",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup([[
            btn("✅ تأكيد الحذف", f"delb_ok_{int(cutoff.timestamp())}"),
            btn("❌ الغاء", "delb_no"),
        ]]),
    )


async def run_delete_before(bot, cutoff):
    docs = [d async for d in db.students.find({"submitted_at": {"$lt": cutoff}})]
    failed = 0
    for d in docs:
        if not await delete_post_msgs(bot, d):
            failed += 1
        await asyncio.sleep(0.05)

    uids = list({d["user_id"] for d in docs})
    await db.students.delete_many({"_id": {"$in": [d["_id"] for d in docs]}})

    # طلبات التعديل المعلقة لهؤلاء
    pend = [p async for p in db.pending_edits.find({"user_id": {"$in": uids}})]
    for p in pend:
        if p.get("request_msg_id"):
            try:
                await bot.delete_message(CHANNEL_ID, p["request_msg_id"])
            except Exception:
                pass
    await db.pending_edits.delete_many({"user_id": {"$in": uids}})

    # نرجع حالتهم للصفر (بدون إشعار لهم)
    await db.users.update_many(
        {"user_id": {"$in": uids}},
        {"$set": {"state": None, "mode": None, "tmp_name": None, "draft": None}},
    )
    return len(docs), failed


async def build_report_text():
    conf = await get_config()
    now_b = datetime.now(BAGHDAD)
    start_utc = now_b.replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc)

    new_users = await db.users.count_documents({"joined_at": {"$gte": start_utc}})
    new_subs = await db.students.count_documents(
        {"submitted_at": {"$gte": start_utc}, "is_admin_post": {"$ne": True}}
    )
    total = await db.students.count_documents({"is_admin_post": {"$ne": True}})
    waiting = len(await pending_user_ids())
    pend = await db.pending_edits.count_documents({})

    if is_closed(conf):
        status = "🔴 متوقف"
    else:
        status = "🟢 يعمل"
        cf = fmt_close(conf)
        if cf:
            status += f" (يغلق تلقائياً {cf})"

    return (
        f"📋 <b>تقرير يوم {now_b.strftime('%Y-%m-%d')}</b>\n\n"
        f"🆕 مستخدمين جدد اليوم: {new_users}\n"
        f"📥 طلبات جديدة اليوم: {new_subs}\n"
        f"🎓 إجمالي الطلاب: {total}\n"
        f"⏰ ضغطوا ستارت وما كملوا: {waiting}\n"
        f"🔄 طلبات تعديل معلقة: {pend}\n"
        f"الاستقبال: {status}"
    )


async def send_report(bot):
    await bot.send_message(CHANNEL_ID, await build_report_text(), parse_mode="HTML")


async def report_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await has_perm(update.effective_user, "report"):
        return
    try:
        await send_report(context.bot)
        await update.message.reply_text("✅ انرسل التقرير للقناة.")
    except Exception as e:
        await update.message.reply_text(f"❌ فشل إرسال التقرير: {e}")


def _safe_cells(ws):
    """يمنع أي اسم يبدأ بـ = من أن يتحول لمعادلة بالإكسل."""
    for cell in ws[ws.max_row]:
        if cell.data_type == "f":
            cell.data_type = "s"


async def do_export(bot, chat_id):
    if Workbook is None:
        await bot.send_message(chat_id, "❌ مكتبة openpyxl غير مثبتة. ضيفها بـ requirements.txt.")
        return

    users = {u["user_id"]: u async for u in db.users.find({}, {"user_id": 1, "username": 1})}

    wb = Workbook()
    ws = wb.active
    ws.title = "الطلاب"
    ws.sheet_view.rightToLeft = True
    ws.append(["#", "الاسم", "الآيدي", "اليوزر", "وقت الإرسال (توقيت العراق)", "نوع المرفق", "ملاحظة"])

    i = 0
    async for s in db.students.find().sort("submitted_at", 1):
        i += 1
        dt = s.get("submitted_at")
        dt_str = ""
        if dt:
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            dt_str = dt.astimezone(BAGHDAD).strftime("%Y-%m-%d %H:%M")
        uname = (users.get(s["user_id"]) or {}).get("username")
        ws.append([
            i,
            s.get("name", ""),
            s["user_id"],
            f"@{uname}" if uname else "",
            dt_str,
            "صورة" if s.get("ftype") == "photo" else "ملف",
            "منشور مشرف" if s.get("is_admin_post") else "",
        ])
        _safe_cells(ws)

    for col, width in zip("ABCDEFG", (6, 32, 14, 20, 26, 12, 14)):
        ws.column_dimensions[col].width = width

    bio = io.BytesIO()
    wb.save(bio)
    stamp = datetime.now(BAGHDAD).strftime("%Y-%m-%d_%H%M")
    await bot.send_document(
        chat_id,
        document=bio.getvalue(),
        filename=f"students_{stamp}.xlsx",
        caption=f"📤 قائمة الطلاب ({i} منشور)",
    )


def _fmt_dt(dt):
    if not dt:
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(BAGHDAD).strftime("%Y-%m-%d %H:%M")


async def _send_xlsx(bot, chat_id, wb, prefix, caption):
    bio = io.BytesIO()
    wb.save(bio)
    stamp = datetime.now(BAGHDAD).strftime("%Y-%m-%d_%H%M")
    await bot.send_document(
        chat_id, document=bio.getvalue(), filename=f"{prefix}_{stamp}.xlsx", caption=caption
    )


async def do_export_pending(bot, chat_id):
    """ملف: اللي ضغطوا ستارت وما كملوا (اسم الحساب + اليوزر + الآيدي + تاريخ الدخول)."""
    if Workbook is None:
        await bot.send_message(chat_id, "❌ مكتبة openpyxl غير مثبتة. ضيفها بـ requirements.txt.")
        return
    ids = set(await pending_user_ids())
    wb = Workbook()
    ws = wb.active
    ws.title = "ما كملوا"
    ws.sheet_view.rightToLeft = True
    ws.append(["#", "اسم الحساب", "اليوزر", "الآيدي", "تاريخ الدخول (توقيت العراق)"])
    i = 0
    async for u in db.users.find({"user_id": {"$in": list(ids)}}).sort("joined_at", 1):
        i += 1
        ws.append([
            i,
            u.get("full_name", "") or "",
            f"@{u['username']}" if u.get("username") else "",
            u["user_id"],
            _fmt_dt(u.get("joined_at")),
        ])
        _safe_cells(ws)
    for col, width in zip("ABCDE", (6, 30, 20, 14, 26)):
        ws.column_dimensions[col].width = width
    await _send_xlsx(bot, chat_id, wb, "not_completed", f"📤 اللي ما كملوا ({i})")


async def do_export_all(bot, chat_id):
    """ملف: كل الطلاب بدون يوزر وآيدي، مع حالة الإكمال."""
    if Workbook is None:
        await bot.send_message(chat_id, "❌ مكتبة openpyxl غير مثبتة. ضيفها بـ requirements.txt.")
        return
    students = {}
    async for st in db.students.find().sort("submitted_at", 1):
        students[st["user_id"]] = st

    wb = Workbook()
    ws = wb.active
    ws.title = "الكل"
    ws.sheet_view.rightToLeft = True
    ws.append(["#", "اسم الحساب", "الاسم المكتوب", "وقت الإرسال (توقيت العراق)", "نوع الملف", "الحالة"])

    rows_done, rows_not = [], []
    async for u in db.users.find().sort("joined_at", 1):
        st = students.pop(u["user_id"], None)
        (rows_done if st else rows_not).append((u, st))
    # منشورات بدون سجل مستخدم (احتياط)
    for uid, st in students.items():
        rows_done.append(({"user_id": uid, "full_name": ""}, st))
    rows_done.sort(key=lambda r: r[1].get("submitted_at") or datetime.min.replace(tzinfo=timezone.utc))

    i = 0
    for u, st in rows_done + rows_not:
        i += 1
        if st:
            ws.append([
                i, u.get("full_name", "") or "", st.get("name", ""),
                _fmt_dt(st.get("submitted_at")),
                "صورة" if st.get("ftype") == "photo" else "ملف",
                "كمل ✅",
            ])
        else:
            ws.append([i, u.get("full_name", "") or "", "", "", "", "ما كمل ❌"])
        _safe_cells(ws)
    for col, width in zip("ABCDEF", (6, 30, 32, 26, 12, 12)):
        ws.column_dimensions[col].width = width
    await _send_xlsx(bot, chat_id, wb, "all_students", f"📤 كل الطلاب ({i})")


async def remove_from_lists(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/remove ID : يحذف الشخص من قوائم البوت فقط. القناة ما تنمس."""
    if not await has_perm(update.effective_user, "delete"):
        return
    if not context.args:
        await update.message.reply_text("مثال: /remove 123456789")
        return
    try:
        target = int(context.args[0])
    except ValueError:
        await update.message.reply_text("❌ الآيدي غير صحيح.")
        return
    if await db.admins.find_one({"user_id": target}):
        await update.message.reply_text("❌ ما يصير تحذف مشرف.")
        return
    u = await db.users.delete_one({"user_id": target})
    st = await db.students.delete_one({"user_id": target})
    await db.pending_edits.delete_many({"user_id": target})
    if not (u.deleted_count or st.deleted_count):
        await update.message.reply_text("ℹ️ ما لقيت هذا الآيدي بالقوائم.")
        return
    await update.message.reply_text(
        f"🧹 انحذف {target} من القوائم.\nمنشوره بالقناة بقى مثل ما هو."
    )


async def export_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await has_perm(update.effective_user, "export"):
        return
    await do_export(context.bot, update.effective_chat.id)


# ============================================================
# لوحة الإدارة
# ============================================================
def fmt_close(conf):
    ca = close_time(conf)
    return ca.astimezone(BAGHDAD).strftime("%Y-%m-%d %H:%M") if ca else None


async def build_panel(user):
    conf = await get_config()
    perms = await perms_of(user)
    owner = is_owner(user)

    closed = is_closed(conf)
    edit_on = conf["edit_enabled"]

    if conf["maintenance_mode"]:
        status = "🔴 متوقف"
    elif closed:
        status = "🔴 مغلق تلقائياً (انتهى الموعد)"
    else:
        status = "🟢 يعمل"

    text = (
        "⚙️ <b>لوحة التحكم</b>\n\n"
        f"الاستقبال: {status}\n"
        f"التعديل: {'🟢 مفعّل' if edit_on else '🔴 متوقف'}"
    )
    if "toggle" in perms:
        cf = fmt_close(conf)
        if cf and not closed:
            text += f"\nالإغلاق التلقائي: ⏰ {cf} (توقيت العراق)"
        elif not cf:
            text += "\nالإغلاق التلقائي: غير محدد"
    if "ping" in perms:
        if conf["ping_enabled"]:
            text += f"\nالبينك: 🟢 كل {conf['ping_interval']} دقيقة"
        else:
            text += "\nالبينك: 🔴 متوقف"
        if not PING_URL:
            text += "\n⚠️ ما كو رابط للبينك (RENDER_EXTERNAL_URL / PING_URL)"
    if "report" in perms:
        if conf["report_enabled"]:
            text += f"\nالتقرير اليومي: 🟢 الساعة {conf['report_time']} (للقناة)"
        else:
            text += "\nالتقرير اليومي: 🔴 متوقف"

    rows = []
    if "stats" in perms:
        rows.append([btn("📊 الإحصائيات", "stats")])
    if "export" in perms:
        rows.append([btn("📤 تصدير المكملين (Excel)", "export")])
        rows.append([btn("⏰ تصدير اللي ما كملوا", "export_pending")])
        rows.append([btn("👥 تصدير الكل", "export_all")])
    if "broadcast" in perms:
        rows.append([btn("📢 إذاعة للجميع", "bc")])
        rows.append([btn("⏰ تذكير اللي ما كملوا", "remind")])
    if "messages" in perms:
        rows.append([btn("📝 رسائل البوت", "msgs")])
    if "toggle" in perms:
        rows.append([btn("▶️ تشغيل الاستقبال" if closed else "⏸ إيقاف الاستقبال", "tg_maint")])
        rows.append([btn("⏸ إيقاف التعديل" if edit_on else "▶️ تشغيل التعديل", "tg_edit")])
        rows.append([btn("⏰ موعد الإغلاق التلقائي", "close_at")])
    if "ping" in perms:
        rows.append([btn("⏸ إيقاف البينك" if conf["ping_enabled"] else "▶️ تشغيل البينك", "tg_ping")])
        rows.append([btn("⏱ تغيير مدة البينك", "ping_int")])
    if "report" in perms:
        rows.append([btn("⏸ إيقاف التقرير اليومي" if conf["report_enabled"] else "▶️ تشغيل التقرير اليومي", "tg_report")])
        rows.append([btn("📋 تقرير الآن (للقناة)", "report_now")])
        rows.append([btn("⏰ وقت التقرير", "report_time")])
    if owner:
        rows.append([btn("👥 المشرفين وصلاحياتهم", "adm_list")])

    if not rows:
        text += "\n\nما عندك صلاحيات حالياً، راجع المطور."
    if "ban" in perms:
        text += "\n\n🚫 الحظر: <code>/ban ID</code> | <code>/unban ID</code>"
    if "delete" in perms:
        text += "\n🗑 حذف طلب: <code>/delete ID</code>"
        text += "\n🧹 حذف كل ما قبل تاريخ: <code>/delete_before 2026-10-20</code>"
        text += "\n🗂 حذف شخص من القوائم فقط: <code>/remove ID</code>"
    if "report" in perms:
        text += "\n📋 <code>/report</code>"
    if "export" in perms:
        text += "\n📤 <code>/export</code>"
    if owner:
        text += "\n👑 <code>/add_admin ID</code> | <code>/rem_admin ID</code>"

    return text, InlineKeyboardMarkup(rows) if rows else None


async def admin_panel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if await is_banned(user.id):
        return
    if not await is_admin(user):
        conf = await get_config()
        if conf["not_admin_msg"]:  # فاضية = يتجاهل بصمت
            await update.message.reply_text(conf["not_admin_msg"])
        return
    await touch_user(user, context)
    text, markup = await build_panel(user)
    await update.message.reply_text(text, reply_markup=markup, parse_mode="HTML")


async def render_msgs_menu(query):
    items = list(MSG_LABELS.items())
    rows = []
    for i in range(0, len(items), 2):
        rows.append([btn(label, f"msg_{key}") for key, label in items[i:i + 2]])
    rows.append([btn("⬅️ رجوع", "panel")])
    await safe_edit(query, "📝 <b>رسائل البوت</b>\nاختر الرسالة اللي تريد تعدلها:", InlineKeyboardMarkup(rows))


async def render_admin_list(query):
    docs = [d async for d in db.admins.find()]
    rows = []
    for d in docs:
        label = d.get("name") or str(d["user_id"])
        rows.append([btn(f"👤 {label}", f"adm_{d['user_id']}")])
    rows.append([btn("⬅️ رجوع", "panel")])
    text = "👥 <b>المشرفين</b>\n" + ("اختر مشرف لضبط صلاحياته:" if docs else "ما كو مشرفين. أضف بـ /add_admin ID")
    await safe_edit(query, text, InlineKeyboardMarkup(rows))


async def render_admin_perms(query, uid):
    doc = await db.admins.find_one({"user_id": uid})
    if not doc:
        await safe_edit(query, "المشرف غير موجود.", InlineKeyboardMarkup([[btn("⬅️ رجوع", "adm_list")]]))
        return
    perms = doc.get("perms", {})
    name = html.escape(doc.get("name") or str(uid))
    rows = [
        [btn(("✅ " if perms.get(k) else "❌ ") + label, f"perm_{uid}_{k}")]
        for k, label in PERMS.items()
    ]
    rows.append([btn("🗑 إزالة المشرف", f"adm_rm_{uid}")])
    rows.append([btn("⬅️ رجوع", "adm_list")])
    text = f"👤 <b>{name}</b>\n🆔 <code>{uid}</code>\n\nاضغط على الصلاحية لتشغيلها أو إيقافها:"
    await safe_edit(query, text, InlineKeyboardMarkup(rows))


async def send_stats(query):
    conf = await get_config()
    stats = await db.command("dbstats")
    data_kb = stats.get("dataSize", 0) / 1024
    size_str = f"{data_kb / 1024:.2f} MB" if data_kb > 1024 else f"{data_kb:.2f} KB"

    users_count = await db.users.count_documents({})
    students_count = await db.students.count_documents({"is_admin_post": {"$ne": True}})
    admin_posts = await db.students.count_documents({"is_admin_post": True})
    pending_count = await db.pending_edits.count_documents({})
    admins_count = await db.admins.count_documents({})
    banned_count = await db.banned.count_documents({})
    waiting = len(await pending_user_ids())

    text = (
        "📊 <b>الإحصائيات:</b>\n\n"
        f"👥 المستخدمين: {users_count}\n"
        f"🎓 الطلاب المسجلين: {students_count}\n"
        f"⏰ ضغطوا ستارت وما كملوا: {waiting}\n"
        f"📌 منشورات المشرفين: {admin_posts}\n"
        f"📥 الطلبات المستلمة: {conf.get('total_received', 0)}\n"
        f"🔄 طلبات تعديل معلقة: {pending_count}\n"
        f"🛡 المشرفين: {admins_count}\n"
        f"🚫 المحظورين: {banned_count}\n"
        f"💾 القاعدة: {size_str} من 512 MB"
    )
    await query.message.reply_text(text, parse_mode="HTML")


async def pending_user_ids():
    """المستخدمين اللي ضغطوا ستارت وما كملوا إرسال الاسم والصورة."""
    submitted = {d["user_id"] async for d in db.students.find({}, {"user_id": 1})}
    banned = {d["user_id"] async for d in db.banned.find({}, {"user_id": 1})}
    admins = {d["user_id"] async for d in db.admins.find({}, {"user_id": 1})}
    conf = await get_config()
    skip = submitted | banned | admins | {conf.get("owner_id")}
    return [u["user_id"] async for u in db.users.find({}, {"user_id": 1}) if u["user_id"] not in skip]


async def copy_to_many(context, ids, msg):
    count = 0
    for uid in ids:
        try:
            await context.bot.copy_message(
                chat_id=uid, from_chat_id=msg.chat_id, message_id=msg.message_id
            )
            count += 1
        except Exception:
            pass
        await asyncio.sleep(0.05)
    return count


# ============================================================
# أزرار القناة (قبول / رفض طلبات التعديل)
# ============================================================
async def resolve_edit(query, context, uid, approve):
    admin = query.from_user
    pend = await db.pending_edits.find_one_and_delete({"user_id": uid})
    if not pend:
        await query.answer("تم البت بهذا الطلب مسبقاً.", show_alert=True)
        return

    who = who_is(admin)
    conf = await get_config()

    if approve:
        student = await db.students.find_one({"user_id": uid})
        try:
            new_id, new_info_id = await replace_post(
                context.bot, uid, student, pend["new_name"], pend["file_id"], pend["ftype"]
            )
        except Exception as e:
            logger.error(f"replace_post failed: {e}")
            # نرجع الطلب حتى ما يضيع
            pend.pop("_id", None)
            await db.pending_edits.update_one({"user_id": uid}, {"$set": pend}, upsert=True)
            await query.answer(f"❌ فشل تعديل المنشور: {e}", show_alert=True)
            return
        await db.students.update_one(
            {"user_id": uid},
            {"$set": {
                "name": pend["new_name"],
                "file_id": pend["file_id"],
                "ftype": pend["ftype"],
                "channel_msg_id": new_id,
                "info_msg_id": new_info_id,
            }},
            upsert=True,
        )
        status = f"✅ وافق عليها: {who}"
        notify_key = "edit_approved_msg"
    else:
        status = f"❌ رفضها: {who}"
        notify_key = "edit_rejected_msg"

    caption = (
        "🔄 <b>طلب تعديل</b>\n"
        f"👤 الاسم القديم: {html.escape(pend.get('old_name', '-'))}\n"
        f"✏️ الاسم الجديد: {html.escape(pend['new_name'])}\n\n"
        f"{status}"
    )
    try:
        await query.edit_message_caption(caption=caption, parse_mode="HTML", reply_markup=None)
    except Exception as e:
        logger.warning(f"edit request caption: {e}")

    await query.answer("تم ✅")

    try:
        await context.bot.send_message(uid, conf[notify_key])
    except Exception as e:
        logger.warning(f"notify student failed: {e}")


async def channel_buttons(query, context, data):
    user = query.from_user
    action, _, rest = data.partition("_")
    try:
        uid = int(rest)
    except ValueError:
        await query.answer()
        return

    if not await is_admin(user):
        await query.answer("⛔ هذا الزر للمشرفين فقط.", show_alert=True)
        return

    if not await has_perm(user, "edits"):
        await query.answer("⛔ ما عندك صلاحية الموافقة على التعديلات.", show_alert=True)
        return
    await resolve_edit(query, context, uid, approve=(action == "ok"))


# ============================================================
# تأكيد الطالب (معاينة قبل الإرسال)
# ============================================================
async def send_preview(context, user, conf, name, file_id, ftype, mode):
    draft = {"name": name, "file_id": file_id, "ftype": ftype}
    caption = f"👤 {html.escape(name)}\n\n{html.escape(conf['preview_msg'])}"
    kb = InlineKeyboardMarkup([[btn("✅ تأكيد", "cf_ok"), btn("🔄 إعادة", "cf_no")]])
    try:
        await send_file(context.bot, user.id, file_id, ftype, caption, kb)
    except Exception as e:
        logger.error(f"preview failed: {e}")
        await say(context, user.id, conf["error_msg"])
        return
    await set_state(user.id, "confirm", mode, name, draft)


async def student_confirm(query, context, data):
    user = query.from_user
    if await is_banned(user.id):
        await query.answer()
        return

    u = await db.users.find_one({"user_id": user.id}) or {}
    draft = u.get("draft")
    if u.get("state") != "confirm" or not draft:
        await query.answer("انتهت صلاحية هذه الرسالة.", show_alert=True)
        try:
            await query.edit_message_reply_markup(reply_markup=None)
        except Exception:
            pass
        return

    conf = await get_config()
    if is_closed(conf) and not await is_admin(user):
        await query.answer(conf["maintenance_msg"], show_alert=True)
        return

    await query.answer()
    try:
        await query.edit_message_reply_markup(reply_markup=None)
    except Exception:
        pass

    mode = u.get("mode") or "new"
    if data == "cf_ok":
        if mode == "edit":
            await submit_edit(context, user, conf, draft["name"], draft["file_id"], draft["ftype"])
        else:
            await submit_new(context, user, conf, draft["name"], draft["file_id"], draft["ftype"])
    else:
        await set_state(user.id, "name", mode)
        await say(context, user.id, conf["edit_ask_name_msg"] if mode == "edit" else conf["ask_name_msg"])


# ============================================================
# معالج الأزرار الرئيسي
# ============================================================
async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    user = query.from_user
    data = query.data or ""

    if data.startswith(("ok_", "no_")):
        await channel_buttons(query, context, data)
        return

    if data in ("cf_ok", "cf_no"):
        await student_confirm(query, context, data)
        return

    await query.answer()

    if not await is_admin(user):
        await safe_edit(query, "لا تملك صلاحية استخدام هذه الأزرار.")
        return

    perms = await perms_of(user)
    owner = is_owner(user)

    async def deny():
        await query.message.reply_text("⛔ ما عندك صلاحية.")

    if data == "panel":
        text, markup = await build_panel(user)
        await safe_edit(query, text, markup)

    elif data == "stats":
        if "stats" not in perms:
            return await deny()
        await send_stats(query)

    elif data == "export":
        if "export" not in perms:
            return await deny()
        await do_export(context.bot, query.message.chat_id)

    elif data == "export_pending":
        if "export" not in perms:
            return await deny()
        await do_export_pending(context.bot, query.message.chat_id)

    elif data == "export_all":
        if "export" not in perms:
            return await deny()
        await do_export_all(context.bot, query.message.chat_id)

    elif data == "bc":
        if "broadcast" not in perms:
            return await deny()
        context.user_data["admin_state"] = "broadcast"
        await query.message.reply_text(
            "📢 <b>وضع الإذاعة</b>\nارسل الرسالة (نص/صورة/ملف) اللي تريد توصل للجميع.\n\n(للإلغاء ارسل: الغاء)",
            parse_mode="HTML",
        )

    elif data == "remind":
        if "broadcast" not in perms:
            return await deny()
        count = len(await pending_user_ids())
        if count == 0:
            await query.message.reply_text("✅ ما كو أحد ضاغط ستارت وما كمل.")
            return
        context.user_data["admin_state"] = "remind"
        await query.message.reply_text(
            f"⏰ <b>تذكير غير المكملين</b>\nعدد المستهدفين: <b>{count}</b>\n"
            "ارسل رسالة التذكير (نص/صورة) وراح توصل لهم فقط.\n\n(للإلغاء ارسل: الغاء)",
            parse_mode="HTML",
        )

    elif data == "msgs":
        if "messages" not in perms:
            return await deny()
        await render_msgs_menu(query)

    elif data.startswith("msg_"):
        if "messages" not in perms:
            return await deny()
        key = data[4:]
        if key not in MSG_LABELS:
            return
        conf = await get_config()
        context.user_data["admin_state"] = f"msg:{key}"
        current = html.escape(conf[key]) if conf[key] else "(فاضية)"
        hint = "\n(ارسل - إذا تريد البوت يتجاهل الأمر بصمت)" if key == "not_admin_msg" else ""
        await query.message.reply_text(
            f"📝 <b>{MSG_LABELS[key]}</b>\n\nالنص الحالي:\n<pre>{current}</pre>\n"
            f"ارسل النص الجديد.{hint}\n(للإلغاء ارسل: الغاء)",
            parse_mode="HTML",
        )

    elif data == "tg_maint":
        if "toggle" not in perms:
            return await deny()
        conf = await get_config()
        if is_closed(conf):
            # تشغيل: نلغي الإيقاف اليدوي وأي موعد إغلاق انتهى
            await set_config("maintenance_mode", False)
            ca = close_time(conf)
            if ca and datetime.now(timezone.utc) >= ca:
                await set_config("close_at", None)
        else:
            await set_config("maintenance_mode", True)
        text, markup = await build_panel(user)
        await safe_edit(query, text, markup)

    elif data == "tg_edit":
        if "toggle" not in perms:
            return await deny()
        conf = await get_config()
        await set_config("edit_enabled", not conf["edit_enabled"])
        text, markup = await build_panel(user)
        await safe_edit(query, text, markup)

    elif data == "close_at":
        if "toggle" not in perms:
            return await deny()
        conf = await get_config()
        cf = fmt_close(conf) or "غير محدد"
        context.user_data["admin_state"] = "close_at"
        await query.message.reply_text(
            f"⏰ <b>موعد الإغلاق التلقائي</b>\nالحالي: {cf}\n\n"
            "ارسل التاريخ والوقت بتوقيت العراق بهذا الشكل:\n<code>2026-10-20 23:59</code>\n"
            "أو ارسل <code>-</code> لإلغاء الموعد.\n\n(للإلغاء ارسل: الغاء)",
            parse_mode="HTML",
        )

    elif data == "tg_ping":
        if "ping" not in perms:
            return await deny()
        conf = await get_config()
        await set_config("ping_enabled", not conf["ping_enabled"])
        text, markup = await build_panel(user)
        await safe_edit(query, text, markup)

    elif data == "ping_int":
        if "ping" not in perms:
            return await deny()
        context.user_data["admin_state"] = "ping_interval"
        await query.message.reply_text(
            "⏱ ارسل المدة بالدقائق (رقم من 1 إلى 14).\nRender يوقف الخدمة بعد 15 دقيقة بدون زيارة.\n\n(للإلغاء ارسل: الغاء)"
        )

    elif data == "tg_report":
        if "report" not in perms:
            return await deny()
        conf = await get_config()
        turning_on = not conf["report_enabled"]
        await set_config("report_enabled", turning_on)
        if turning_on:
            # إذا وقت التقرير مرّ اليوم، أول تقرير يكون باجر (مو فوراً)
            now_b = datetime.now(BAGHDAD)
            if now_b.strftime("%H:%M") >= conf["report_time"]:
                await set_config("report_last_date", now_b.strftime("%Y-%m-%d"))
        text, markup = await build_panel(user)
        await safe_edit(query, text, markup)

    elif data == "report_now":
        if "report" not in perms:
            return await deny()
        try:
            await send_report(context.bot)
            await query.message.reply_text("✅ انرسل التقرير للقناة.")
        except Exception as e:
            await query.message.reply_text(f"❌ فشل إرسال التقرير: {e}")

    elif data == "report_time":
        if "report" not in perms:
            return await deny()
        conf = await get_config()
        context.user_data["admin_state"] = "report_time"
        await query.message.reply_text(
            f"⏰ <b>وقت التقرير اليومي</b>\nالحالي: {conf['report_time']} (توقيت العراق)\n\n"
            "ارسل الوقت الجديد بهذا الشكل: <code>22:00</code>\n\n(للإلغاء ارسل: الغاء)",
            parse_mode="HTML",
        )

    elif data == "delb_no":
        await safe_edit(query, "✅ تم إلغاء الحذف.")

    elif data.startswith("delb_ok_"):
        if "delete" not in perms:
            return await deny()
        try:
            cutoff = datetime.fromtimestamp(int(data[8:]), timezone.utc)
        except ValueError:
            return
        await safe_edit(query, "⏳ جاري الحذف، انتظر...")

        async def job():
            try:
                n, failed = await run_delete_before(context.bot, cutoff)
                result = f"✅ تم حذف {n} منشور."
                if failed:
                    result += f"\n⚠️ فشل حذف {failed} منشور من القناة، احذفهم يدوياً."
            except Exception as e:
                logger.error(f"delete_before failed: {e}")
                result = f"❌ صار خطأ: {html.escape(str(e))}"
            await safe_edit(query, result)

        context.application.create_task(job())

    elif data == "adm_list":
        if not owner:
            return await deny()
        await render_admin_list(query)

    elif data.startswith("adm_rm_"):
        if not owner:
            return await deny()
        try:
            uid = int(data[7:])
        except ValueError:
            return
        await db.admins.delete_one({"user_id": uid})
        await render_admin_list(query)

    elif data.startswith("adm_"):
        if not owner:
            return await deny()
        try:
            uid = int(data[4:])
        except ValueError:
            return
        await render_admin_perms(query, uid)

    elif data.startswith("perm_"):
        if not owner:
            return await deny()
        try:
            _, uid_s, perm = data.split("_", 2)
            uid = int(uid_s)
        except ValueError:
            return
        if perm not in PERMS:
            return
        doc = await db.admins.find_one({"user_id": uid})
        if not doc:
            return
        new_val = not doc.get("perms", {}).get(perm, False)
        await db.admins.update_one({"user_id": uid}, {"$set": {f"perms.{perm}": new_val}})
        await render_admin_perms(query, uid)


# ============================================================
# معالج الرسائل (خاص فقط)
# ============================================================
async def handle_admin_state(update, context, user, state):
    msg = update.message

    if is_cancel(msg.text):
        context.user_data["admin_state"] = None
        await msg.reply_text("✅ تم إلغاء العملية.")
        return

    if state in ("broadcast", "remind"):
        if not await has_perm(user, "broadcast"):
            context.user_data["admin_state"] = None
            await msg.reply_text("⛔ ما عندك صلاحية.")
            return
        await msg.reply_text("⏳ جاري الإرسال...")
        if state == "broadcast":
            banned_ids = {d["user_id"] async for d in db.banned.find({}, {"user_id": 1})}
            ids = [u["user_id"] async for u in db.users.find({}, {"user_id": 1}) if u["user_id"] not in banned_ids]
        else:
            ids = await pending_user_ids()
        count = await copy_to_many(context, ids, msg)
        context.user_data["admin_state"] = None
        await msg.reply_text(f"✅ تم الإرسال لـ {count} مستخدم.")

    elif state.startswith("msg:"):
        if not await has_perm(user, "messages"):
            context.user_data["admin_state"] = None
            await msg.reply_text("⛔ ما عندك صلاحية تعديل الرسائل.")
            return
        key = state[4:]
        if key not in MSG_LABELS:
            context.user_data["admin_state"] = None
            return
        if not msg.text:
            await msg.reply_text("❌ ارسل نص فقط.")
            return
        new_text = msg.text
        if key == "not_admin_msg" and new_text.strip() == "-":
            new_text = ""
        await set_config(key, new_text)
        context.user_data["admin_state"] = None
        await msg.reply_text("✅ تم تغيير الرسالة، تظهر للطلاب فوراً." if new_text else "✅ صار البوت يتجاهل الأمر بصمت.")

    elif state == "ping_interval":
        if not await has_perm(user, "ping"):
            context.user_data["admin_state"] = None
            await msg.reply_text("⛔ ما عندك صلاحية.")
            return
        try:
            n = int((msg.text or "").strip())
        except ValueError:
            n = 0
        if not 1 <= n <= 14:
            await msg.reply_text("❌ ارسل رقم من 1 إلى 14.")
            return
        await set_config("ping_interval", n)
        context.user_data["admin_state"] = None
        await msg.reply_text(f"✅ صار البينك كل {n} دقيقة.")

    elif state == "report_time":
        if not await has_perm(user, "report"):
            context.user_data["admin_state"] = None
            await msg.reply_text("⛔ ما عندك صلاحية.")
            return
        try:
            t = datetime.strptime((msg.text or "").strip(), "%H:%M")
        except ValueError:
            await msg.reply_text("❌ الصيغة غلط. مثال: 22:00")
            return
        hhmm = t.strftime("%H:%M")
        await set_config("report_time", hhmm)
        now_b = datetime.now(BAGHDAD)
        if now_b.strftime("%H:%M") >= hhmm:
            await set_config("report_last_date", now_b.strftime("%Y-%m-%d"))
        context.user_data["admin_state"] = None
        await msg.reply_text(f"✅ صار التقرير اليومي الساعة {hhmm} (توقيت العراق).")

    elif state == "close_at":
        if not await has_perm(user, "toggle"):
            context.user_data["admin_state"] = None
            await msg.reply_text("⛔ ما عندك صلاحية.")
            return
        text = (msg.text or "").strip()
        if text == "-":
            await set_config("close_at", None)
            context.user_data["admin_state"] = None
            await msg.reply_text("✅ تم إلغاء موعد الإغلاق التلقائي.")
            return
        try:
            dt = datetime.strptime(text, "%Y-%m-%d %H:%M").replace(tzinfo=BAGHDAD)
        except ValueError:
            await msg.reply_text("❌ الصيغة غلط. مثال: 2026-10-20 23:59")
            return
        if dt <= datetime.now(timezone.utc):
            await msg.reply_text("❌ الموعد لازم يكون بالمستقبل.")
            return
        await set_config("close_at", dt.astimezone(timezone.utc))
        context.user_data["admin_state"] = None
        await msg.reply_text(f"✅ راح يتوقف الاستقبال تلقائياً بتاريخ {text} (توقيت العراق).")


async def submit_new(context, user, conf, name, file_id, ftype, admin=False):
    try:
        sent = await send_post(context.bot, user.id, name, file_id, ftype)
    except Exception as e:
        logger.error(f"channel post failed: {e}")
        await notify_owner_error(
            context, f"⚠️ فشل النشر بالقناة: {e}\nتأكد من CHANNEL_ID وأن البوت مشرف بالقناة."
        )
        # نرجعه لخطوة الصورة حتى يكدر يعيد الإرسال
        await set_state(user.id, "photo", "new", name)
        await say(context, user.id, conf["error_msg"])
        return

    info_id = await send_info(context.bot, user.id, user.full_name, user.username, sent.message_id)
    doc = {
        "user_id": user.id,
        "name": name,
        "file_id": file_id,
        "ftype": ftype,
        "channel_msg_id": sent.message_id,
        "info_msg_id": info_id,
        "submitted_at": datetime.now(timezone.utc),
    }
    if admin:
        # كل إرسال من المشرف منشور مستقل
        await db.students.insert_one({**doc, "is_admin_post": True})
    else:
        await db.students.update_one({"user_id": user.id}, {"$set": doc}, upsert=True)

    await set_state(user.id, None)
    await inc_received()
    await say(context, user.id, with_edit_note(conf, "success_msg") if not admin else conf["success_msg"])


async def submit_edit(context, user, conf, name, file_id, ftype):
    if not conf["edit_enabled"]:
        await set_state(user.id, None)
        await say(context, user.id, conf["edit_disabled_msg"])
        return
    if await db.pending_edits.find_one({"user_id": user.id}):
        await set_state(user.id, None)
        await say(context, user.id, conf["edit_pending_msg"])
        return

    student = await db.students.find_one({"user_id": user.id}) or {}
    old_name = student.get("name", "-")

    try:
        req = await send_request(context.bot, user.id, old_name, name, file_id, ftype)
    except Exception as e:
        logger.error(f"edit request failed: {e}")
        await notify_owner_error(context, f"⚠️ فشل إرسال طلب تعديل للقناة: {e}")
        await set_state(user.id, "photo", "edit", name)
        await say(context, user.id, conf["error_msg"])
        return

    await db.pending_edits.update_one(
        {"user_id": user.id},
        {"$set": {
            "user_id": user.id,
            "old_name": old_name,
            "new_name": name,
            "file_id": file_id,
            "ftype": ftype,
            "request_msg_id": req.message_id,
        }},
        upsert=True,
    )
    await set_state(user.id, None)
    await say(context, user.id, conf["edit_sent_msg"])


async def handle_student(update, context, user, conf):
    msg = update.message
    admin = await is_admin(user)
    u = await db.users.find_one({"user_id": user.id}) or {}
    state = u.get("state")
    mode = u.get("mode") or "new"
    student = None if admin else await db.students.find_one({"user_id": user.id})
    text = (msg.text or "").strip()

    # إلغاء التعديل
    if state and mode == "edit" and is_cancel(text):
        await set_state(user.id, None)
        await msg.reply_text("تم الإلغاء ✅")
        return

    if state == "confirm":
        await msg.reply_text(conf["confirm_hint_msg"])
        return

    if state == "name":
        if not msg.text or not valid_name(text):
            await msg.reply_text(conf["bad_name_msg"])
            return
        await set_state(user.id, "photo", mode, text)
        await msg.reply_text(conf["edit_ask_photo_msg"] if mode == "edit" else conf["ask_photo_msg"])
        return

    if state == "photo":
        file_id, ftype = extract_image(msg)
        if not file_id:
            await msg.reply_text(conf["wrong_type_msg"])
            return
        name = u.get("tmp_name")
        if not name:
            await set_state(user.id, "name", mode)
            await msg.reply_text(conf["edit_ask_name_msg"] if mode == "edit" else conf["ask_name_msg"])
            return
        if admin:
            # المشرف ينشر مباشرة بدون معاينة
            await submit_new(context, user, conf, name, file_id, ftype, admin=True)
        else:
            await send_preview(context, user, conf, name, file_id, ftype, mode)
        return

    # ما كو خطوة نشطة
    if admin:
        # المشرف يكدر يرسل أكثر من مرة: أي رسالة تبدأ دور جديد
        if msg.text and valid_name(text):
            await set_state(user.id, "photo", "new", text)
            await msg.reply_text(conf["ask_photo_msg"])
        else:
            await set_state(user.id, "name", "new")
            await msg.reply_text(conf["ask_name_msg"])
        return

    if msg.text and text == "تعديل":
        if not conf["edit_enabled"]:
            await msg.reply_text(conf["edit_disabled_msg"])
        elif not student:
            await set_state(user.id, "name", "new")
            await msg.reply_text(conf["ask_name_msg"])
        elif await db.pending_edits.find_one({"user_id": user.id}):
            await msg.reply_text(conf["edit_pending_msg"])
        else:
            await set_state(user.id, "name", "edit")
            await msg.reply_text(conf["edit_ask_name_msg"])
        return

    if student:
        await msg.reply_text(with_edit_note(conf, "already_msg"))
        return

    await set_state(user.id, "name", "new")
    await msg.reply_text(conf["ask_name_msg"])


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    msg = update.message
    if not msg or await is_banned(user.id):
        return

    await touch_user(user, context)

    admin_state = context.user_data.get("admin_state")
    if admin_state and await is_admin(user):
        await handle_admin_state(update, context, user, admin_state)
        return

    conf = await get_config()
    if is_closed(conf) and not await is_admin(user):
        await msg.reply_text(conf["maintenance_msg"])
        return

    await handle_student(update, context, user, conf)


async def start_word(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """كلمة «ستارت» بالعربي تشتغل مثل /start (إلا إذا المشرف وسط إدخال نص للوحة)."""
    if context.user_data.get("admin_state"):
        await handle_message(update, context)
    else:
        await start(update, context)


# ============================================================
# البينك التلقائي (حتى ما ينام البوت على Render) + سيرفر صحة بسيط
# ============================================================
class _Health(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-Type", "text/plain")
        self.end_headers()
        self.wfile.write(b"OK")

    def do_HEAD(self):
        self.send_response(200)
        self.end_headers()

    def log_message(self, *args):
        pass


def start_health_server():
    server = ThreadingHTTPServer(("0.0.0.0", PORT), _Health)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    logger.info(f"health server on port {PORT}")


async def ping_loop():
    """يزور رابط الخدمة نفسه حسب الإعدادات المخزنة بالقاعدة (تشغيل/إيقاف + المدة)."""
    last = None
    async with httpx.AsyncClient(timeout=20) as http:
        while True:
            try:
                conf = await get_config()
                interval = max(1, int(conf.get("ping_interval", 10))) * 60
                due = last is None or (time.monotonic() - last) >= interval
                if conf.get("ping_enabled") and PING_URL and due:
                    last = time.monotonic()
                    try:
                        r = await http.get(PING_URL)
                        logger.info(f"ping {PING_URL} -> {r.status_code}")
                    except Exception as e:
                        logger.warning(f"ping failed: {e}")
            except Exception as e:
                logger.warning(f"ping loop error: {e}")
            await asyncio.sleep(20)


async def report_loop(bot):
    """يرسل التقرير اليومي للقناة بالوقت المحدد (مرة وحدة باليوم) إذا مفعّل."""
    while True:
        try:
            conf = await get_config()
            if conf.get("report_enabled"):
                now_b = datetime.now(BAGHDAD)
                today = now_b.strftime("%Y-%m-%d")
                if now_b.strftime("%H:%M") >= conf.get("report_time", "22:00") and conf.get("report_last_date") != today:
                    await set_config("report_last_date", today)  # قبل الإرسال حتى ما يتكرر إذا فشل
                    try:
                        await send_report(bot)
                    except Exception as e:
                        logger.error(f"daily report failed: {e}")
                        owner_id = conf.get("owner_id")
                        if owner_id:
                            try:
                                await bot.send_message(owner_id, f"⚠️ فشل إرسال التقرير اليومي للقناة: {e}")
                            except Exception:
                                pass
        except Exception as e:
            logger.warning(f"report loop error: {e}")
        await asyncio.sleep(60)


async def post_init(application):
    application.bot_data["tasks"] = [
        asyncio.create_task(ping_loop()),
        asyncio.create_task(report_loop(application.bot)),
    ]


async def post_shutdown(application):
    for task in application.bot_data.get("tasks", []):
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass


# ============================================================
# التشغيل
# ============================================================
def main():
    if not BOT_TOKEN or not MONGO_URI or CHANNEL_ID is None:
        print("❌ خطأ: تأكد من إضافة المتغيرات BOT_TOKEN و MONGO_URI و CHANNEL_ID (رقم صحيح).")
        return

    start_health_server()
    application = Application.builder().token(BOT_TOKEN).post_init(post_init).post_shutdown(post_shutdown).build()

    private = filters.ChatType.PRIVATE
    application.add_handler(CommandHandler("start", start, filters=private))
    application.add_handler(CommandHandler("help", help_command, filters=private))
    application.add_handler(CommandHandler("admin", admin_panel, filters=private))
    application.add_handler(CommandHandler("add_admin", add_admin, filters=private))
    application.add_handler(CommandHandler("rem_admin", rem_admin, filters=private))
    application.add_handler(CommandHandler("ban", ban_user, filters=private))
    application.add_handler(CommandHandler("unban", unban_user, filters=private))
    application.add_handler(CommandHandler("delete", delete_request, filters=private))
    application.add_handler(CommandHandler("export", export_command, filters=private))
    application.add_handler(CommandHandler("remove", remove_from_lists, filters=private))
    application.add_handler(CommandHandler("delete_before", delete_before_cmd, filters=private))
    application.add_handler(CommandHandler("report", report_command, filters=private))
    application.add_handler(CallbackQueryHandler(button_handler))
    application.add_handler(MessageHandler(private & filters.Regex(r"^\s*ستارت\s*$"), start_word))
    application.add_handler(MessageHandler(private & ~filters.COMMAND, handle_message))

    print("🤖 البوت يعمل...", flush=True)
    application.run_polling()


if __name__ == "__main__":
    main()
